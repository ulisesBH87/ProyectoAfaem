import io
from datetime import datetime, time
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.modelos.consumo_modelos import BitacoraConsumo, CatalogoTarifas
from app.modelos.usuario_modelo import Usuario
from app.modelos.equipo_modelo import Equipos
from app.modelos.catalogos_liga_modelo import Ligas

class ReporteExcelServicio:
    @staticmethod
    def sort_report_rows(rows):
        """
        Ordena las filas descendentemente:
        1. costo_total_mxn descendente
        2. costo_total_usd descendente
        3. operaciones totales descendente
        """
        def sort_key(x):
            costo_mxn = float(x.get("costo_total_mxn") or x.get("costo_mxn") or 0.0)
            costo_usd = float(x.get("costo_total_usd") or x.get("costo_usd") or 0.0)
            ocr = int(x.get("ocr_count") or x.get("ocr") or 0)
            foto = int(x.get("foto_count") or x.get("foto") or 0)
            vm = int(x.get("verificamex_count") or x.get("vm") or 0)
            total_ops = ocr + foto + vm
            return (-costo_mxn, -costo_usd, -total_ops)
            
        return sorted(rows, key=sort_key)

    @staticmethod
    def normalizar_tipo_servicio(tipo_raw: str) -> str:
        if not tipo_raw:
            return ""
        t = str(tipo_raw).strip().upper()
        if t in ("PHOTO SCAN", "PHOTO_SCAN", "ESCANEO FOTOGRAFÍA", "ESCANEO FOTOGRAFIA", "FOTOGRAFÍA", "FOTOGRAFIA"):
            return "PHOTO_SCAN"
        if t in ("VERIFICAMEX", "VERIFICA MEX", "VERIFICACIÓN CURP", "VERIFICACION CURP", "VERIFICACIÓN CURP/CIUDADANO", "VERIFICACION CURP/CIUDADANO"):
            return "VERIFICAMEX"
        if t in ("OCR", "ESCANEO OCR"):
            return "OCR"
        return tipo_raw

    @staticmethod
    def _parse_fecha_inicio(fecha_inicio: str | None):
        if not fecha_inicio or fecha_inicio in ("todos", "Inicio"):
            return None
        try:
            return datetime.fromisoformat(fecha_inicio)
        except ValueError:
            return None

    @staticmethod
    def _parse_fecha_fin_inclusiva(fecha_fin: str | None):
        if not fecha_fin or fecha_fin in ("todos", "Fin"):
            return None
        try:
            fecha = datetime.fromisoformat(fecha_fin)
            if len(str(fecha_fin)) <= 10:
                return datetime.combine(fecha.date(), time.max)
            return fecha
        except ValueError:
            return None

    @staticmethod
    def formatear_fecha_reporte(fecha_str: str) -> str:
        if not fecha_str or fecha_str in ("Inicio", "Fin", "todos"):
            return fecha_str
        try:
            dt = datetime.strptime(fecha_str.strip(), "%Y-%m-%d")
            return dt.strftime("%d/%m/%Y")
        except Exception:
            return fecha_str

    @classmethod
    def generar_reporte_consumos(
        cls,
        db,
        resumen: dict | None,
        auditoria: dict | None,
        fecha_inicio: str,
        fecha_fin: str,
        tipo_reporte: str = "todos",
        liga_id: int | None = None,
        equipo_id: int | None = None
    ) -> bytes:
        # 1. Consultar y filtrar registros del Ledger de consumos
        query = db.query(BitacoraConsumo)
        
        fecha_inicio_dt = cls._parse_fecha_inicio(fecha_inicio)
        fecha_fin_dt = cls._parse_fecha_fin_inclusiva(fecha_fin)
        
        if fecha_inicio_dt:
            query = query.filter(BitacoraConsumo.CreadoEn >= fecha_inicio_dt)
        if fecha_fin_dt:
            query = query.filter(BitacoraConsumo.CreadoEn <= fecha_fin_dt)
            
        if liga_id:
            query = query.filter(BitacoraConsumo.LigaId == liga_id)
        if equipo_id:
            query = query.filter(BitacoraConsumo.EquipoId == equipo_id)
            
        if tipo_reporte in ("jugador", "jugadores"):
            query = query.filter(BitacoraConsumo.TipoRegistro.notin_(["PRESIDENTE", "ENTRENADOR"]))
        elif tipo_reporte in ("directivo", "directivos"):
            query = query.filter(BitacoraConsumo.TipoRegistro.in_(["PRESIDENTE", "ENTRENADOR"]))
            
        registros = query.all()
        
        # 2. Cargar mapas de catálogo para resolver nombres
        usuarios_db = db.query(Usuario).all()
        usuarios_map = {}
        for u in usuarios_db:
            nombre_usr = "SISTEMA / INVITADO"
            rol_usr = "INVITADO"
            if u.PersonaRelacion:
                p = u.PersonaRelacion
                nombre_usr = f"{p.Nombre} {p.PrimerApellido} {p.SegundoApellido or ''}".strip().upper()
            if u.RolRelacion:
                rol_usr = u.RolRelacion.Nombre.upper()
            usuarios_map[u.UsuarioId] = {"nombre": nombre_usr, "rol": rol_usr}
            
        equipos_db = db.query(Equipos).all()
        equipos_map = {e.EquipoId: e.NombreEquipo.upper() for e in equipos_db}
        
        ligas_db = db.query(Ligas).all()
        ligas_map = {l.LigaId: l.Nombreliga.upper() for l in ligas_db}
        
        # 3. Procesar agrupaciones e indicadores en memoria
        total_operaciones = len(registros)
        costo_total_usd = 0.0
        costo_total_mxn = 0.0
        
        operaciones_por_servicio = {}
        costo_por_operacion = {}
        
        desglose_jugadores = {}
        desglose_directivos = {}
        desglose_equipos = {}
        desglose_ligas = {}
        
        for r in registros:
            tipo_con = r.TipoConsumo
            costo_val = float(r.CostoTotal)
            divisa = r.Divisa or "MXN"
            
            if divisa == "USD":
                costo_total_usd += costo_val
            else:
                costo_total_mxn += costo_val
                
            operaciones_por_servicio[tipo_con] = operaciones_por_servicio.get(tipo_con, 0) + 1
            costo_por_operacion[tipo_con] = costo_por_operacion.get(tipo_con, 0.0) + costo_val
            
            # Auditoría
            usr_info = usuarios_map.get(r.UsuarioId, {"nombre": "INVITADO", "rol": "INVITADO"})
            ejecutor_nombre = usr_info["nombre"]
            ejecutor_rol = usr_info["rol"]
            
            jug_nombre = r.JugadorNombre or "SIN NOMBRE (EN OCR)"
            jug_curp = r.JugadorCURP or "SIN CURP"
            
            eq_id = r.EquipoId
            eq_nombre = equipos_map.get(eq_id, "PRE-REGISTRO / SIN EQUIPO") if eq_id else "PRE-REGISTRO / SIN EQUIPO"
            
            lg_id = r.LigaId
            lg_nombre = ligas_map.get(lg_id, "SIN LIGA") if lg_id else "SIN LIGA"
            
            costo_usd = costo_val if divisa == "USD" else 0.0
            costo_mxn = costo_val if divisa == "MXN" else 0.0
            
            # Agrupar por Jugador o Directivo
            if r.TipoRegistro in ("PRESIDENTE", "ENTRENADOR"):
                dir_key = (r.UsuarioId, jug_nombre, jug_curp, eq_id, lg_id, r.TipoRegistro)
                if dir_key not in desglose_directivos:
                    desglose_directivos[dir_key] = {
                        "ejecutor_nombre": ejecutor_nombre,
                        "ejecutor_rol": ejecutor_rol,
                        "directivo_nombre": jug_nombre,
                        "directivo_curp": jug_curp,
                        "directivo_rol": r.TipoRegistro,
                        "equipo_nombre": eq_nombre,
                        "liga_nombre": lg_nombre,
                        "ocr_count": 0,
                        "foto_count": 0,
                        "verificamex_count": 0,
                        "costo_total_usd": 0.0,
                        "costo_total_mxn": 0.0
                    }
                item_dir = desglose_directivos[dir_key]
                if tipo_con == "OCR":
                    item_dir["ocr_count"] += 1
                elif tipo_con == "PHOTO_SCAN":
                    item_dir["foto_count"] += 1
                elif tipo_con == "VERIFICAMEX":
                    item_dir["verificamex_count"] += 1
                item_dir["costo_total_usd"] += costo_usd
                item_dir["costo_total_mxn"] += costo_mxn
            else:
                jug_key = (r.UsuarioId, jug_nombre, jug_curp, eq_id, lg_id)
                if jug_key not in desglose_jugadores:
                    desglose_jugadores[jug_key] = {
                        "ejecutor_nombre": ejecutor_nombre,
                        "ejecutor_rol": ejecutor_rol,
                        "jugador_nombre": jug_nombre,
                        "jugador_curp": jug_curp,
                        "equipo_nombre": eq_nombre,
                        "liga_nombre": lg_nombre,
                        "ocr_count": 0,
                        "foto_count": 0,
                        "verificamex_count": 0,
                        "costo_total_usd": 0.0,
                        "costo_total_mxn": 0.0
                    }
                item_jug = desglose_jugadores[jug_key]
                if tipo_con == "OCR":
                    item_jug["ocr_count"] += 1
                elif tipo_con == "PHOTO_SCAN":
                    item_jug["foto_count"] += 1
                elif tipo_con == "VERIFICAMEX":
                    item_jug["verificamex_count"] += 1
                item_jug["costo_total_usd"] += costo_usd
                item_jug["costo_total_mxn"] += costo_mxn
                
            if eq_id:
                if eq_id not in desglose_equipos:
                    desglose_equipos[eq_id] = {
                        "equipo_id": eq_id,
                        "equipo_nombre": eq_nombre,
                        "liga_nombre": lg_nombre,
                        "ocr_count": 0,
                        "foto_count": 0,
                        "verificamex_count": 0,
                        "costo_total_usd": 0.0,
                        "costo_total_mxn": 0.0
                    }
                item_eq = desglose_equipos[eq_id]
                if tipo_con == "OCR":
                    item_eq["ocr_count"] += 1
                elif tipo_con == "PHOTO_SCAN":
                    item_eq["foto_count"] += 1
                elif tipo_con == "VERIFICAMEX":
                    item_eq["verificamex_count"] += 1
                item_eq["costo_total_usd"] += costo_usd
                item_eq["costo_total_mxn"] += costo_mxn
                
            if lg_id:
                if lg_id not in desglose_ligas:
                    desglose_ligas[lg_id] = {
                        "liga_id": lg_id,
                        "liga_nombre": lg_nombre,
                        "ocr_count": 0,
                        "foto_count": 0,
                        "verificamex_count": 0,
                        "costo_total_usd": 0.0,
                        "costo_total_mxn": 0.0
                    }
                item_lg = desglose_ligas[lg_id]
                if tipo_con == "OCR":
                    item_lg["ocr_count"] += 1
                elif tipo_con == "PHOTO_SCAN":
                    item_lg["foto_count"] += 1
                elif tipo_con == "VERIFICAMEX":
                    item_lg["verificamex_count"] += 1
                item_lg["costo_total_usd"] += costo_usd
                item_lg["costo_total_mxn"] += costo_mxn

        # 4. Crear libro de Excel
        wb = openpyxl.Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)
        
        # Estilos Generales
        title_font = Font(name="Segoe UI", size=16, bold=True, color="000000")
        sheet_title_font = Font(name="Segoe UI", size=14, bold=True, color="000000")
        subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="475569")
        section_font = Font(name="Segoe UI", size=12, bold=True, color="1E293B")
        
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        cell_font = Font(name="Segoe UI", size=10, color="0F172A")
        
        thin_border_side = Side(style='thin', color='E2E8F0')
        border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        # Formatos de número de Excel
        MXN_FORMAT = '"$"#,##0.00'
        USD_FORMAT = '"$"#,##0.0000'
        INT_FORMAT = '#,##0'
        
        def style_header_row(ws, row_idx, max_col):
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[row_idx].height = 28

        def finalize_sheet(ws, start_row, end_row, max_col):
            ws.views.sheetView[0].showGridLines = True
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    val_str = str(cell.value or '')
                    if cell.number_format and ('$' in cell.number_format or '#' in cell.number_format):
                        val_str = "$999,999.00"
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 13)
            
            if end_row > start_row:
                last_col_letter = get_column_letter(max_col)
                ws.auto_filter.ref = f"A{start_row}:{last_col_letter}{end_row}"

        # =====================================================================
        # 1. HOJA: Resumen
        # =====================================================================
        ws_res = wb.create_sheet(title="Resumen")
        ws_res.sheet_properties.tabColor = "000000"
        
        ws_res["A1"] = "REPORTE DE CONSUMO DE SERVICIOS"
        ws_res["A1"].font = title_font
        
        # Formatear fechas en Día/Mes/Año
        f_ini_fmt = cls.formatear_fecha_reporte(fecha_inicio)
        f_fin_fmt = cls.formatear_fecha_reporte(fecha_fin)
        ws_res["A2"] = f"Periodo: {f_ini_fmt} al {f_fin_fmt}"
        ws_res["A2"].font = subtitle_font
        
        ws_res["A3"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        ws_res["A3"].font = subtitle_font
        
        # Filtros informativos adicionales en el Resumen
        extra_filtro_text = []
        if tipo_reporte != "todos":
            extra_filtro_text.append(f"Categoría: {tipo_reporte.upper()}")
        if liga_id:
            extra_filtro_text.append(f"Liga: {ligas_map.get(liga_id, 'DESCONOCIDA')}")
        if equipo_id:
            extra_filtro_text.append(f"Equipo: {equipos_map.get(equipo_id, 'DESCONOCIDO')}")
        if extra_filtro_text:
            ws_res["A4"] = f"Filtros aplicados: {', '.join(extra_filtro_text)}"
            ws_res["A4"].font = Font(name="Segoe UI", size=9, bold=True, color="475569")
            
        # KPIs
        ws_res["A6"] = "Métricas Generales"
        ws_res["A6"].font = section_font
        
        ws_res["A7"] = "Total de Operaciones"
        ws_res["B7"] = total_operaciones
        ws_res["B7"].number_format = INT_FORMAT
        
        ws_res["A8"] = "Costo Total USD"
        ws_res["B8"] = costo_total_usd
        ws_res["B8"].number_format = USD_FORMAT
        
        ws_res["A9"] = "Costo Total MXN"
        ws_res["B9"] = costo_total_mxn
        ws_res["B9"].number_format = MXN_FORMAT
        
        for r in range(7, 10):
            ws_res.cell(row=r, column=1).font = Font(name="Segoe UI", size=10, bold=True, color="475569")
            ws_res.cell(row=r, column=2).font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
            ws_res.cell(row=r, column=1).border = border_all
            ws_res.cell(row=r, column=2).border = border_all
            ws_res.cell(row=r, column=1).fill = zebra_fill
            ws_res.cell(row=r, column=2).alignment = Alignment(horizontal="right")
            ws_res.row_dimensions[r].height = 20
            
        # Tabla resumen de servicios
        ws_res["A12"] = "Costo por Servicio"
        ws_res["A12"].font = section_font
        
        headers_res = ["Servicio", "Proveedor", "Operaciones", "Costo Acumulado", "Divisa", "Descripción"]
        for col_idx, h in enumerate(headers_res, 1):
            ws_res.cell(row=13, column=col_idx, value=h)
        style_header_row(ws_res, 13, len(headers_res))
        
        # Agrupar tarifas activas
        tarifas_db = db.query(CatalogoTarifas).filter(CatalogoTarifas.Estatus == True).all()
        grouped_services = {}
        for t in tarifas_db:
            norm = cls.normalizar_tipo_servicio(t.TipoConsumo)
            if not norm:
                continue
            if norm not in grouped_services:
                grouped_services[norm] = {
                    "tipo": norm,
                    "proveedor": t.Proveedor or "DEFAULT",
                    "cantidad": 0,
                    "costo": 0.0,
                    "divisa": t.Divisa or "MXN",
                    "descripcion": t.Descripcion or ""
                }
                
        # Completar con operaciones y costos del periodo
        for tipo_raw, cant in operaciones_por_servicio.items():
            norm = cls.normalizar_tipo_servicio(tipo_raw)
            if norm not in grouped_services:
                grouped_services[norm] = {
                    "tipo": norm,
                    "proveedor": "DEFAULT",
                    "cantidad": cant,
                    "costo": 0.0,
                    "divisa": "MXN",
                    "descripcion": ""
                }
            else:
                grouped_services[norm]["cantidad"] = cant
                
        for tipo_raw, cost in costo_por_operacion.items():
            norm = cls.normalizar_tipo_servicio(tipo_raw)
            if norm in grouped_services:
                grouped_services[norm]["costo"] = float(cost or 0.0)
                
        known_services = list(grouped_services.values())
        service_order = ['OCR', 'PHOTO_SCAN', 'VERIFICAMEX']
        known_services.sort(key=lambda x: service_order.index(x["tipo"]) if x["tipo"] in service_order else 999)
        
        current_row = 14
        for s in known_services:
            tipo = s["tipo"]
            proveedor = s["proveedor"]
            if tipo == "PHOTO_SCAN" and str(proveedor).upper() == "DEFAULT":
                proveedor = "PHOTO SCAN"
                
            display_name = "Escaneo OCR" if tipo == "OCR" else ("Escaneo Fotografía" if tipo == "PHOTO_SCAN" else ("Verificación CURP" if tipo == "VERIFICAMEX" else tipo))
            
            ws_res.cell(row=current_row, column=1, value=display_name)
            ws_res.cell(row=current_row, column=2, value=proveedor)
            
            c_cant = ws_res.cell(row=current_row, column=3, value=s["cantidad"])
            c_cant.number_format = INT_FORMAT
            c_cant.alignment = Alignment(horizontal="right")
            
            divisa = s["divisa"]
            c_cost = ws_res.cell(row=current_row, column=4, value=s["costo"])
            c_cost.number_format = USD_FORMAT if divisa == "USD" else MXN_FORMAT
            c_cost.alignment = Alignment(horizontal="right")
            
            ws_res.cell(row=current_row, column=5, value=divisa)
            ws_res.cell(row=current_row, column=6, value=s["descripcion"])
            
            for c_idx in range(1, 7):
                cell = ws_res.cell(row=current_row, column=c_idx)
                cell.font = cell_font
                cell.border = border_all
                if current_row % 2 == 0:
                    cell.fill = zebra_fill
            ws_res.row_dimensions[current_row].height = 20
            current_row += 1
            
        finalize_sheet(ws_res, start_row=13, end_row=current_row-1, max_col=6)
        
        # =====================================================================
        # 2. HOJA: Liga
        # =====================================================================
        if tipo_reporte in ("todos", "liga"):
            ws_liga = wb.create_sheet(title="Liga")
            ws_liga.sheet_properties.tabColor = "000000"
            
            ws_liga["A1"] = "CONSUMOS POR LIGA"
            ws_liga["A1"].font = sheet_title_font
            ws_liga["A2"] = f"Periodo: {f_ini_fmt} al {f_fin_fmt}"
            ws_liga["A2"].font = subtitle_font
            
            headers_liga = ["Liga", "Conteo OCR", "Conteo Foto", "Conteo VerificaMEX", "Costo Total USD", "Costo Total MXN"]
            for col_idx, h in enumerate(headers_liga, 1):
                ws_liga.cell(row=4, column=col_idx, value=h)
            style_header_row(ws_liga, 4, len(headers_liga))
            
            ligas_sorted = cls.sort_report_rows(list(desglose_ligas.values()))
            
            current_row = 5
            for item in ligas_sorted:
                ws_liga.cell(row=current_row, column=1, value=item.get("liga_nombre") or "SIN LIGA")
                
                c_ocr = ws_liga.cell(row=current_row, column=2, value=item.get("ocr_count", 0))
                c_ocr.number_format = INT_FORMAT
                c_ocr.alignment = Alignment(horizontal="right")
                
                c_foto = ws_liga.cell(row=current_row, column=3, value=item.get("foto_count", 0))
                c_foto.number_format = INT_FORMAT
                c_foto.alignment = Alignment(horizontal="right")
                
                c_vm = ws_liga.cell(row=current_row, column=4, value=item.get("verificamex_count", 0))
                c_vm.number_format = INT_FORMAT
                c_vm.alignment = Alignment(horizontal="right")
                
                c_usd = ws_liga.cell(row=current_row, column=5, value=item.get("costo_total_usd", 0.0))
                c_usd.number_format = USD_FORMAT
                c_usd.alignment = Alignment(horizontal="right")
                
                c_mxn = ws_liga.cell(row=current_row, column=6, value=item.get("costo_total_mxn", 0.0))
                c_mxn.number_format = MXN_FORMAT
                c_mxn.alignment = Alignment(horizontal="right")
                
                for c_idx in range(1, 7):
                    cell = ws_liga.cell(row=current_row, column=c_idx)
                    cell.font = cell_font
                    cell.border = border_all
                    if current_row % 2 == 0:
                        cell.fill = zebra_fill
                ws_liga.row_dimensions[current_row].height = 20
                current_row += 1
                
            finalize_sheet(ws_liga, start_row=4, end_row=current_row-1, max_col=6)
            
        # =====================================================================
        # 3. HOJA: Equipo
        # =====================================================================
        if tipo_reporte in ("todos", "equipo"):
            ws_eq = wb.create_sheet(title="Equipo")
            ws_eq.sheet_properties.tabColor = "000000"
            
            ws_eq["A1"] = "CONSUMOS POR EQUIPO"
            ws_eq["A1"].font = sheet_title_font
            ws_eq["A2"] = f"Periodo: {f_ini_fmt} al {f_fin_fmt}"
            ws_eq["A2"].font = subtitle_font
            
            headers_eq = ["Equipo", "Liga", "Conteo OCR", "Conteo Foto", "Conteo VerificaMEX", "Costo Total USD", "Costo Total MXN"]
            for col_idx, h in enumerate(headers_eq, 1):
                ws_eq.cell(row=4, column=col_idx, value=h)
            style_header_row(ws_eq, 4, len(headers_eq))
            
            equipos_sorted = cls.sort_report_rows(list(desglose_equipos.values()))
            
            current_row = 5
            for item in equipos_sorted:
                ws_eq.cell(row=current_row, column=1, value=item.get("equipo_nombre") or "SIN EQUIPO")
                ws_eq.cell(row=current_row, column=2, value=item.get("liga_nombre") or "SIN LIGA")
                
                c_ocr = ws_eq.cell(row=current_row, column=3, value=item.get("ocr_count", 0))
                c_ocr.number_format = INT_FORMAT
                c_ocr.alignment = Alignment(horizontal="right")
                
                c_foto = ws_eq.cell(row=current_row, column=4, value=item.get("foto_count", 0))
                c_foto.number_format = INT_FORMAT
                c_foto.alignment = Alignment(horizontal="right")
                
                c_vm = ws_eq.cell(row=current_row, column=5, value=item.get("verificamex_count", 0))
                c_vm.number_format = INT_FORMAT
                c_vm.alignment = Alignment(horizontal="right")
                
                c_usd = ws_eq.cell(row=current_row, column=6, value=item.get("costo_total_usd", 0.0))
                c_usd.number_format = USD_FORMAT
                c_usd.alignment = Alignment(horizontal="right")
                
                c_mxn = ws_eq.cell(row=current_row, column=7, value=item.get("costo_total_mxn", 0.0))
                c_mxn.number_format = MXN_FORMAT
                c_mxn.alignment = Alignment(horizontal="right")
                
                for c_idx in range(1, 8):
                    cell = ws_eq.cell(row=current_row, column=c_idx)
                    cell.font = cell_font
                    cell.border = border_all
                    if current_row % 2 == 0:
                        cell.fill = zebra_fill
                ws_eq.row_dimensions[current_row].height = 20
                current_row += 1
                
            finalize_sheet(ws_eq, start_row=4, end_row=current_row-1, max_col=7)
            
        # =====================================================================
        # 4. HOJA: Jugadores
        # =====================================================================
        if tipo_reporte in ("todos", "jugador", "jugadores"):
            ws_jug = wb.create_sheet(title="Jugadores")
            ws_jug.sheet_properties.tabColor = "000000"
            
            ws_jug["A1"] = "CONSUMOS POR JUGADORES"
            ws_jug["A1"].font = sheet_title_font
            ws_jug["A2"] = f"Periodo: {f_ini_fmt} al {f_fin_fmt}"
            ws_jug["A2"].font = subtitle_font
            
            headers_jug = ["Nombre Jugador", "Equipo", "Liga", "Conteo OCR", "Conteo Foto", "Conteo VerificaMEX", "Costo Total USD", "Costo Total MXN"]
            for col_idx, h in enumerate(headers_jug, 1):
                ws_jug.cell(row=4, column=col_idx, value=h)
            style_header_row(ws_jug, 4, len(headers_jug))
            
            jugadores_sorted = cls.sort_report_rows(list(desglose_jugadores.values()))
            
            current_row = 5
            for item in jugadores_sorted:
                ws_jug.cell(row=current_row, column=1, value=item.get("jugador_nombre") or "SIN NOMBRE")
                ws_jug.cell(row=current_row, column=2, value=item.get("equipo_nombre") or "PRE-REGISTRO / SIN EQUIPO")
                ws_jug.cell(row=current_row, column=3, value=item.get("liga_nombre") or "SIN LIGA")
                
                c_ocr = ws_jug.cell(row=current_row, column=4, value=item.get("ocr_count", 0))
                c_ocr.number_format = INT_FORMAT
                c_ocr.alignment = Alignment(horizontal="right")
                
                c_foto = ws_jug.cell(row=current_row, column=5, value=item.get("foto_count", 0))
                c_foto.number_format = INT_FORMAT
                c_foto.alignment = Alignment(horizontal="right")
                
                c_vm = ws_jug.cell(row=current_row, column=6, value=item.get("verificamex_count", 0))
                c_vm.number_format = INT_FORMAT
                c_vm.alignment = Alignment(horizontal="right")
                
                c_usd = ws_jug.cell(row=current_row, column=7, value=item.get("costo_total_usd", 0.0))
                c_usd.number_format = USD_FORMAT
                c_usd.alignment = Alignment(horizontal="right")
                
                c_mxn = ws_jug.cell(row=current_row, column=8, value=item.get("costo_total_mxn", 0.0))
                c_mxn.number_format = MXN_FORMAT
                c_mxn.alignment = Alignment(horizontal="right")
                
                for c_idx in range(1, 9):
                    cell = ws_jug.cell(row=current_row, column=c_idx)
                    cell.font = cell_font
                    cell.border = border_all
                    if current_row % 2 == 0:
                        cell.fill = zebra_fill
                ws_jug.row_dimensions[current_row].height = 20
                current_row += 1
                
            finalize_sheet(ws_jug, start_row=4, end_row=current_row-1, max_col=8)
            
        # =====================================================================
        # 5. HOJA: Directivos
        # =====================================================================
        if tipo_reporte in ("todos", "directivo", "directivos"):
            ws_dir = wb.create_sheet(title="Directivos")
            ws_dir.sheet_properties.tabColor = "000000"
            
            ws_dir["A1"] = "CONSUMOS POR DIRECTIVOS"
            ws_dir["A1"].font = sheet_title_font
            ws_dir["A2"] = f"Periodo: {f_ini_fmt} al {f_fin_fmt}"
            ws_dir["A2"].font = subtitle_font
            
            headers_dir = ["Nombre Directivo", "Rol", "Equipo", "Liga", "Conteo OCR", "Conteo Foto", "Conteo VerificaMEX", "Costo Total USD", "Costo Total MXN"]
            for col_idx, h in enumerate(headers_dir, 1):
                ws_dir.cell(row=4, column=col_idx, value=h)
            style_header_row(ws_dir, 4, len(headers_dir))
            
            directivos_sorted = cls.sort_report_rows(list(desglose_directivos.values()))
            
            current_row = 5
            for item in directivos_sorted:
                ws_dir.cell(row=current_row, column=1, value=item.get("directivo_nombre") or "SIN NOMBRE")
                ws_dir.cell(row=current_row, column=2, value=item.get("directivo_rol") or "DIRECTIVO")
                ws_dir.cell(row=current_row, column=3, value=item.get("equipo_nombre") or "PRE-REGISTRO / SIN EQUIPO")
                ws_dir.cell(row=current_row, column=4, value=item.get("liga_nombre") or "SIN LIGA")
                
                c_ocr = ws_dir.cell(row=current_row, column=5, value=item.get("ocr_count", 0))
                c_ocr.number_format = INT_FORMAT
                c_ocr.alignment = Alignment(horizontal="right")
                
                c_foto = ws_dir.cell(row=current_row, column=6, value=item.get("foto_count", 0))
                c_foto.number_format = INT_FORMAT
                c_foto.alignment = Alignment(horizontal="right")
                
                c_vm = ws_dir.cell(row=current_row, column=7, value=item.get("verificamex_count", 0))
                c_vm.number_format = INT_FORMAT
                c_vm.alignment = Alignment(horizontal="right")
                
                c_usd = ws_dir.cell(row=current_row, column=8, value=item.get("costo_total_usd", 0.0))
                c_usd.number_format = USD_FORMAT
                c_usd.alignment = Alignment(horizontal="right")
                
                c_mxn = ws_dir.cell(row=current_row, column=9, value=item.get("costo_total_mxn", 0.0))
                c_mxn.number_format = MXN_FORMAT
                c_mxn.alignment = Alignment(horizontal="right")
                
                for c_idx in range(1, 10):
                    cell = ws_dir.cell(row=current_row, column=c_idx)
                    cell.font = cell_font
                    cell.border = border_all
                    if current_row % 2 == 0:
                        cell.fill = zebra_fill
                ws_dir.row_dimensions[current_row].height = 20
                current_row += 1
                
            finalize_sheet(ws_dir, start_row=4, end_row=current_row-1, max_col=9)
            
        # Guardar en memoria y retornar bytes
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
